# Pemakaian: python -I tools/build_pembangkit_kit.py "<PEMBANGKIT_MM_YYYY (FULL).xlsx>" data/pembangkit_kit_YYYY_MM.js "<nama file asli>"
# Lalu ganti PEMBANGKIT_KIT_SRC di index.html ke file baru. Butuh: pip install openpyxl
# Konversi PEMBANGKIT_<MM>_<YYYY> (FULL).xlsx -> data/pembangkit_kit_<YYYY>_<MM>.js
# Sheet2 = 1 baris per titik sentral (jenis kolom C, x/y kolom K/L).
# Sheet1 = 1 baris per mesin (data operasi bulanan) — digabung lewat kode mesin.
# Kolom harga/biaya (H_*, OA_*, NBB_*, COST CENTER) SENGAJA tidak diekspor (repo publik).
import openpyxl,sys,json,collections,datetime
src,out=sys.argv[1],sys.argv[2]
wb=openpyxl.load_workbook(src,read_only=True,data_only=True)
s1=list(wb['Sheet1'].iter_rows(values_only=True)); h=list(s1[0]); R=s1[1:]; ix={c:i for i,c in enumerate(h)}
s2=list(wb['Sheet2'].iter_rows(values_only=True)); h2=list(s2[0]); R2=s2[1:]; i2={c:i for i,c in enumerate(h2)}
bulan,tahun=R[0][0],R[0][1]

T=['']; TI={'':0}  # indeks 0 = kosong, supaya nilai hilang bisa ditulis 0
def t(v):
    v=clean(v)
    if v is None: return None
    if v not in TI: TI[v]=len(T); T.append(v)
    return TI[v]
def clean(v):
    if v is None: return None
    if isinstance(v,float) and v.is_integer(): v=int(v)
    v=str(v).strip()
    return None if v in ('','-','#N/A','None','nan') else v
def num(v,d=0):
    try: f=float(v or 0)
    except: return 0
    return int(round(f)) if d==0 else round(f,d)

FUELS=['HSD','MFO','IDO','BIO','OLN','BBR','GAS','BGS','PBM','AIR','SRY','BMS','LLN','HFO','BSD','CPO','BFM']
PFUELS=['HSD','MFO','IDO','BIO','OLN','BBR','GAS','BGS','BMS','LLN','HFO','BSD','CPO','BFM']
m1={}
for r in R:
    g=lambda c: r[ix[c]]
    m={}
    def put(k,v):
        if v not in (None,0,''): m[k]=v
    put('k',clean(g('KODE MESIN'))); put('n',clean(g('MESIN')))
    put('mk',t(g('MERK'))); put('ty',t(g('TYPE'))); put('sr',clean(g('SERI')))
    put('gn',t(g('MERK GENERATOR'))); put('tr',t(g('NAMA TRAFO STEPUP PEMBANGKIT')))
    put('tg',t(g('TEGANGAN HV/LV (KV)'))); put('kva',clean(g('KAPASITAS (KVA)')))
    put('th',clean(g('TAHUN OPERASI')))
    put('sm',t(g('STATUS MILIK'))); put('smd',t(g('S_MILIK_DETAIL')))
    put('bb',t(g('KODE JENIS BAHAN BAKAR'))); put('jt',t(g('JENIS TEGANGAN')))
    put('bp',num(g('BEBAN PUNCAK (kW)'),1)); put('dt',num(g('DAYA TERPASANG (kW)'),1))
    put('dmn',num(g('DAYA MAMPU NETTO (kW)'),1)); put('dmp',num(g('DAYA MAMPU PASOK (kW)'),1))
    put('mu',t(g('MUTASI')))  # KONDISI sengaja tidak diekspor (keputusan user 2026-10-15)
    for k,c in [('op','JAM_OP (Jam)'),('foh','JAM_FOH (Jam)'),('ph','JAM_PH (Jam)'),('moh','JAM_MOH (Jam)'),('poh','JAM_POH'),('meh','JAM_MEH'),('peh','JAM_PEH'),('efdh','EFDH'),('emdh','EMDH'),('epdh','EPDH')]:
        put(k,num(g(c),2))
    # SdOF (jumlah gangguan) sengaja tidak diekspor (keputusan user 2026-10-15)
    kwh={f:num(g('KWH_'+f)) for f in FUELS if num(g('KWH_'+f))>0}
    if kwh: m['kwh']=kwh
    pk={f:num(g('P_'+f)) for f in PFUELS if num(g('P_'+f))>0}
    if pk: m['pk']=pk
    put('ktt',num(g('KTT_HSD')))
    put('pl',num(g('P_PELUMAS'),1))
    put('opr',t(g('OPERATOR PEMBANGKIT')))
    if m.get('k'): m1[m['k']]=m   # baris tanpa kode mesin dilewati

MF=['k','n','mk','ty','sr','gn','tr','tg','kva','th','sm','smd','bb','jt','bp','dt','dmn','dmp','mu','op','foh','ph','moh','poh','meh','peh','efdh','emdh','epdh','kwh','pk','ktt','pl','opr']
def pos(m):
    a=[m.get(f) for f in MF]
    while a and a[-1] is None: a.pop()
    return [0 if v is None else v for v in a]
S=[];bad=[]
for r in R2:
    g=lambda c: r[i2[c]]
    x,y=g('x'),g('y')
    ok=isinstance(x,(int,float)) and isinstance(y,(int,float)) and 94<x<142 and -12<y<7
    if not ok: bad.append([clean(g('sentral')),clean(g('jenis_pembangkit')),str(x),str(y)])
    ms=[]
    for kd in str(clean(g('kode_mesin')) or '').split('\n'):   # clean(): float bulat -> int, sama dengan kunci m1
        kd=clean(kd) or ''
        if kd in m1: ms.append(pos(m1[kd]))
    s={'k':clean(g('kode_sentral')),'n':clean(g('sentral')),'j':clean(g('jenis_pembangkit')),
       'lat':round(y,6) if ok else None,'lon':round(x,6) if ok else None,
       'reg':t(g('regional')),'wil':t(g('wilayah')),'cab':t(g('cabang')),'rk':clean(g('kode_ranting')),
       'ran':t(g('ranting')),'sis':t(g('sistem')),'prov':t(g('provinsi')),'m':ms}
    S.append({k:v for k,v in s.items() if v is not None})
used=sum(len(s['m']) for s in S)
meta={'periode':'%04d-%02d'%(tahun,bulan),'bulan':bulan,'tahun':tahun,'sumber':sys.argv[3],'mf':MF,
      'dibuat':datetime.date.today().isoformat(),'n_titik':len(S),'n_mesin':used,
      'mesin_tanpa_titik':len(m1)-used,'titik_tanpa_koordinat':bad}
js=('// Data inventaris pembangkit PLN — dihasilkan otomatis dari Excel "%s".\n'
    '// Jangan edit manual: jalankan ulang konverter. Kolom harga/biaya sengaja tidak diekspor.\n'
    'window.PEMBANGKIT_KIT = {"meta":%s,"T":%s,"s":%s};\n')%(meta['sumber'],json.dumps(meta,ensure_ascii=False),
    json.dumps(T,ensure_ascii=False,separators=(',',':')),json.dumps(S,ensure_ascii=False,separators=(',',':')))
open(out,'w',encoding='utf-8').write(js)
print(len(js)/1e6,'MB',meta['n_titik'],used,meta['mesin_tanpa_titik'],bad,len(T))
